# -*- coding: utf-8 -*-
"""按飞书 open_id 隔离的签核统计存储与 Excel 导出。"""

from __future__ import annotations

import datetime as _dt
import hashlib
import io
import math
import sqlite3
from contextlib import closing
from pathlib import Path


SCHEMA = """
CREATE TABLE IF NOT EXISTS sign_actions (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    open_id TEXT NOT NULL,
    display_name TEXT NOT NULL DEFAULT '',
    application_no TEXT NOT NULL DEFAULT '',
    applicant TEXT NOT NULL DEFAULT '',
    description TEXT NOT NULL DEFAULT '',
    uom TEXT NOT NULL DEFAULT '',
    item_type TEXT NOT NULL DEFAULT '',
    action TEXT NOT NULL,
    source TEXT NOT NULL,
    rule_name TEXT NOT NULL DEFAULT '',
    reason TEXT NOT NULL DEFAULT '',
    notify_sent INTEGER NOT NULL DEFAULT 0,
    status TEXT NOT NULL DEFAULT 'verified',
    created_at TEXT NOT NULL
);
CREATE INDEX IF NOT EXISTS idx_sign_actions_user_time
ON sign_actions(open_id, created_at DESC);
CREATE TABLE IF NOT EXISTS processed_events (
    event_key TEXT PRIMARY KEY,
    expires_at REAL NOT NULL
);
CREATE TABLE IF NOT EXISTS work_items (
    open_id TEXT NOT NULL,
    item_key TEXT NOT NULL,
    application_no TEXT NOT NULL DEFAULT '',
    initial_route TEXT NOT NULL,
    current_status TEXT NOT NULL,
    auto_attempts INTEGER NOT NULL DEFAULT 0,
    auto_failures INTEGER NOT NULL DEFAULT 0,
    first_seen_at TEXT NOT NULL,
    last_seen_at TEXT NOT NULL,
    resolved_at TEXT NOT NULL DEFAULT '',
    resolution_source TEXT NOT NULL DEFAULT '',
    duration_ms INTEGER NOT NULL DEFAULT 0,
    PRIMARY KEY (open_id, item_key)
);
CREATE INDEX IF NOT EXISTS idx_work_items_first_seen
ON work_items(first_seen_at, open_id);
CREATE TABLE IF NOT EXISTS run_metrics (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    open_id TEXT NOT NULL,
    trigger TEXT NOT NULL,
    discovered_count INTEGER NOT NULL DEFAULT 0,
    manual_pending_count INTEGER NOT NULL DEFAULT 0,
    auto_success_count INTEGER NOT NULL DEFAULT 0,
    failure_count INTEGER NOT NULL DEFAULT 0,
    duration_ms INTEGER NOT NULL DEFAULT 0,
    created_at TEXT NOT NULL
);
CREATE INDEX IF NOT EXISTS idx_run_metrics_created
ON run_metrics(created_at, open_id);
CREATE TABLE IF NOT EXISTS request_metrics (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    endpoint TEXT NOT NULL,
    status_code INTEGER NOT NULL,
    duration_ms INTEGER NOT NULL,
    created_at TEXT NOT NULL
);
CREATE INDEX IF NOT EXISTS idx_request_metrics_created
ON request_metrics(created_at, endpoint);
"""


def _connect(db_path):
    raw_path = str(db_path)
    if raw_path.startswith("file:"):
        conn = sqlite3.connect(raw_path, timeout=15, uri=True)
        conn.row_factory = sqlite3.Row
        conn.execute("PRAGMA busy_timeout=15000")
        conn.executescript(SCHEMA)
        return conn
    path = Path(db_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(str(path), timeout=15)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA busy_timeout=15000")
    conn.execute("PRAGMA journal_mode=WAL")
    conn.execute("PRAGMA synchronous=NORMAL")
    conn.executescript(SCHEMA)
    return conn


def _now() -> str:
    return _dt.datetime.now().astimezone().isoformat(timespec="milliseconds")


def _item_key(item: dict) -> str:
    application_no = str(item.get("no", "") or "").strip()
    material = application_no or "\x1f".join(
        str(item.get(key, "") or "").strip()
        for key in ("applicant", "description", "desc", "uom", "item_type")
    )
    return hashlib.sha256(material.encode("utf-8")).hexdigest()


def observe_work_items(db_path, open_id: str, items: list[dict]) -> None:
    """记录唯一待办生命周期；不保存额外敏感字段，只保留单号和哈希键。"""
    now = _now()
    with closing(_connect(db_path)) as conn, conn:
        for item in items:
            route = "auto" if item.get("action") in ("approve", "reject") else "manual"
            item_key = _item_key(item)
            current_status = "pending_auto" if route == "auto" else "pending_manual"
            # CentOS 7 常见的 SQLite 3.7.x 不支持新版 UPSERT。先原子插入，再更新可变字段，
            # 保留首次路由/首次发现时间，同时兼容旧 SQLite。
            conn.execute(
                """INSERT OR IGNORE INTO work_items
                (open_id,item_key,application_no,initial_route,current_status,first_seen_at,last_seen_at)
                VALUES (?,?,?,?,?,?,?)""",
                (
                    open_id, item_key, str(item.get("no", "") or ""),
                    route, current_status,
                    now, now,
                ),
            )
            conn.execute(
                """UPDATE work_items SET
                   last_seen_at=?,
                   current_status=CASE
                     WHEN resolved_at<>'' THEN current_status
                     ELSE ? END
                   WHERE open_id=? AND item_key=?""",
                (now, current_status, open_id, item_key),
            )


def record_auto_outcomes(
    db_path, open_id: str, attempted_items: list[dict], successful_nos: set[str]
) -> None:
    """记录自动动作尝试；平台复查未成功的尝试计为失败。"""
    with closing(_connect(db_path)) as conn, conn:
        for item in attempted_items:
            succeeded = str(item.get("no", "") or "") in successful_nos
            conn.execute(
                """UPDATE work_items SET
                   auto_attempts=auto_attempts+1,
                   auto_failures=auto_failures+?,
                   current_status=CASE WHEN ? THEN current_status ELSE 'failed_auto' END
                   WHERE open_id=? AND item_key=?""",
                (0 if succeeded else 1, 1 if succeeded else 0, open_id, _item_key(item)),
            )


def _mark_work_item_resolved(
    conn, open_id: str, item: dict, source: str, resolved_at: str
) -> None:
    key = _item_key(item)
    row = conn.execute(
        "SELECT first_seen_at FROM work_items WHERE open_id=? AND item_key=?",
        (open_id, key),
    ).fetchone()
    if not row:
        route = "auto" if source == "auto" else "manual"
        conn.execute(
            """INSERT INTO work_items
            (open_id,item_key,application_no,initial_route,current_status,
             first_seen_at,last_seen_at)
            VALUES (?,?,?,?,?,?,?)""",
            (
                open_id, key, str(item.get("no", "") or ""), route,
                "pending_" + route, resolved_at, resolved_at,
            ),
        )
        first_seen = resolved_at
    else:
        first_seen = str(row["first_seen_at"])
    try:
        started = _dt.datetime.fromisoformat(first_seen)
        ended = _dt.datetime.fromisoformat(resolved_at)
        duration_ms = max(0, int((ended - started).total_seconds() * 1000))
    except ValueError:
        duration_ms = 0
    conn.execute(
        """UPDATE work_items SET current_status='resolved', resolved_at=?,
           resolution_source=?, duration_ms=?, last_seen_at=?
           WHERE open_id=? AND item_key=?""",
        (resolved_at, source, duration_ms, resolved_at, open_id, key),
    )


def record_action(db_path, open_id: str, display_name: str, item: dict, action: str,
                  source: str, rule_name: str = "", reason: str = "",
                  notify_sent: bool = False, status: str = "verified") -> int:
    created_at = _now()
    values = (
        open_id, display_name, item.get("no", ""), item.get("applicant", ""),
        item.get("description", item.get("desc", "")), item.get("uom", ""),
        item.get("item_type", ""), action, source, rule_name, reason,
        1 if notify_sent else 0, status, created_at,
    )
    with closing(_connect(db_path)) as conn, conn:
        cursor = conn.execute(
            """INSERT INTO sign_actions
            (open_id,display_name,application_no,applicant,description,uom,item_type,
             action,source,rule_name,reason,notify_sent,status,created_at)
            VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)""", values,
        )
        _mark_work_item_resolved(conn, open_id, item, source, created_at)
        return int(cursor.lastrowid)


def record_run_metric(
    db_path, open_id: str, trigger: str, discovered_count: int,
    manual_pending_count: int, auto_success_count: int, failure_count: int,
    duration_ms: int,
) -> None:
    with closing(_connect(db_path)) as conn, conn:
        conn.execute(
            """INSERT INTO run_metrics
            (open_id,trigger,discovered_count,manual_pending_count,
             auto_success_count,failure_count,duration_ms,created_at)
            VALUES (?,?,?,?,?,?,?,?)""",
            (
                open_id, trigger, max(0, int(discovered_count)),
                max(0, int(manual_pending_count)), max(0, int(auto_success_count)),
                max(0, int(failure_count)), max(0, int(duration_ms)), _now(),
            ),
        )


def record_request_metric(
    db_path, endpoint: str, status_code: int, duration_ms: int
) -> None:
    try:
        with closing(_connect(db_path)) as conn, conn:
            conn.execute(
                """INSERT INTO request_metrics(endpoint,status_code,duration_ms,created_at)
                VALUES (?,?,?,?)""",
                (str(endpoint or "unknown")[:120], int(status_code), max(0, int(duration_ms)), _now()),
            )
            conn.execute(
                "DELETE FROM request_metrics WHERE created_at < ?",
                (_window_start(30),),
            )
    except sqlite3.Error:
        # 指标采集不得影响业务响应。
        return


def _window_start(days: int) -> str:
    value = _dt.datetime.now().astimezone() - _dt.timedelta(days=max(1, int(days)))
    return value.isoformat(timespec="seconds")


def production_kpis(db_path, days: int = 7) -> dict:
    """全局生产 KPI；仅供服务端确认过的管理员路由使用。"""
    since = _window_start(days)
    with closing(_connect(db_path)) as conn, conn:
        row = conn.execute(
            """SELECT COUNT(*) AS discovered,
               SUM(CASE WHEN resolution_source='auto' THEN 1 ELSE 0 END) AS automatic,
               SUM(CASE WHEN initial_route='manual' THEN 1 ELSE 0 END) AS manual_route,
               SUM(auto_attempts) AS attempts,
               SUM(auto_failures) AS failures,
               AVG(CASE WHEN resolved_at<>'' THEN duration_ms END) AS avg_duration_ms,
               SUM(CASE WHEN current_status='pending_manual' THEN 1 ELSE 0 END) AS manual_backlog
               FROM work_items WHERE first_seen_at>=?""",
            (since,),
        ).fetchone()
        run_row = conn.execute(
            """SELECT SUM(auto_success_count) AS successes,
               SUM(failure_count) AS failures
               FROM run_metrics WHERE created_at>=?""",
            (since,),
        ).fetchone()
        daily = [
            dict(item) for item in conn.execute(
                """SELECT substr(first_seen_at,1,10) AS day,
                   COUNT(*) AS discovered,
                   SUM(CASE WHEN resolution_source='auto' THEN 1 ELSE 0 END) AS automatic,
                   SUM(CASE WHEN initial_route='manual' THEN 1 ELSE 0 END) AS manual_route,
                   SUM(auto_failures) AS failures
                   FROM work_items WHERE first_seen_at>=?
                   GROUP BY day ORDER BY day""",
                (since,),
            ).fetchall()
        ]
    discovered = int(row["discovered"] or 0)
    attempts = int(row["attempts"] or 0)
    automatic = int(row["automatic"] or 0)
    manual_route = int(row["manual_route"] or 0)
    failures = int(run_row["failures"] or 0)
    successful_attempts = int(run_row["successes"] or 0)
    return {
        "days": max(1, int(days)),
        "discovered": discovered,
        "automatic": automatic,
        "manual_route": manual_route,
        "failures": failures,
        "manual_backlog": int(row["manual_backlog"] or 0),
        "automatic_rate": round(automatic * 100 / discovered, 2) if discovered else 0.0,
        "manual_rate": round(manual_route * 100 / discovered, 2) if discovered else 0.0,
        "failure_rate": round(
            failures * 100 / (successful_attempts + failures), 2
        ) if successful_attempts + failures else 0.0,
        "avg_duration_ms": round(float(row["avg_duration_ms"] or 0), 1),
        "daily": daily,
    }


def _percentile(values: list[int], percentile: float) -> int:
    if not values:
        return 0
    ordered = sorted(values)
    index = min(len(ordered) - 1, max(0, math.ceil(len(ordered) * percentile) - 1))
    return int(ordered[index])


def load_profile(db_path, days: int = 7) -> dict:
    since = _window_start(days)
    with closing(_connect(db_path)) as conn, conn:
        requests = conn.execute(
            "SELECT status_code,duration_ms,created_at FROM request_metrics WHERE created_at>=?",
            (since,),
        ).fetchall()
        runs = conn.execute(
            "SELECT duration_ms,failure_count,created_at FROM run_metrics WHERE created_at>=?",
            (since,),
        ).fetchall()
    per_minute: dict[str, int] = {}
    for row in requests:
        minute = str(row["created_at"])[:16]
        per_minute[minute] = per_minute.get(minute, 0) + 1
    request_durations = [int(row["duration_ms"]) for row in requests]
    run_durations = [int(row["duration_ms"]) for row in runs]
    runs_per_minute: dict[str, int] = {}
    for row in runs:
        minute = str(row["created_at"])[:16]
        runs_per_minute[minute] = runs_per_minute.get(minute, 0) + 1
    return {
        "days": max(1, int(days)),
        "request_count": len(requests),
        "peak_requests_per_minute": max(per_minute.values() or [0]),
        "request_p95_ms": _percentile(request_durations, 0.95),
        "request_error_rate": round(
            sum(1 for row in requests if int(row["status_code"]) >= 500) * 100 / len(requests), 2
        ) if requests else 0.0,
        "run_count": len(runs),
        "peak_runs_per_minute": max(runs_per_minute.values() or [0]),
        "run_p95_ms": _percentile(run_durations, 0.95),
        "run_failure_count": sum(int(row["failure_count"]) for row in runs),
        "has_real_load": bool(requests or runs),
    }


def claim_event(db_path, event_key: str, ttl_seconds: int = 600) -> bool:
    """跨线程/进程认领飞书事件；TTL 内重复事件返回 False。"""
    import time
    key = str(event_key or "").strip()
    if not key:
        return True
    now = time.time()
    with closing(_connect(db_path)) as conn, conn:
        conn.execute("DELETE FROM processed_events WHERE expires_at <= ?", (now,))
        cursor = conn.execute(
            "INSERT OR IGNORE INTO processed_events(event_key, expires_at) VALUES (?, ?)",
            (key, now + max(1, int(ttl_seconds))),
        )
        return cursor.rowcount == 1


def _filters(open_id: str, params: dict):
    where = ["open_id = ?"]
    values = [open_id]
    mapping = {
        "action": "action", "source": "source", "rule": "rule_name",
        "applicant": "applicant", "item_type": "item_type", "uom": "uom",
    }
    for key, column in mapping.items():
        raw_values = params.get(key, [])
        if not isinstance(raw_values, (list, tuple, set)):
            raw_values = [raw_values]
        selected = []
        for raw_value in raw_values:
            value = str(raw_value or "").strip()
            if value and value not in selected:
                selected.append(value)
        if selected:
            placeholders = ",".join("?" for _ in selected)
            where.append("%s IN (%s)" % (column, placeholders))
            values.extend(selected)
    keyword = str(params.get("keyword", "") or "").strip()
    if keyword:
        where.append("(application_no LIKE ? OR applicant LIKE ? OR description LIKE ? OR rule_name LIKE ?)")
        values.extend(["%%%s%%" % keyword] * 4)
    date_from = str(params.get("date_from", "") or "").strip()
    date_to = str(params.get("date_to", "") or "").strip()
    if date_from:
        where.append("created_at >= ?")
        values.append(date_from + "T00:00:00")
    if date_to:
        where.append("created_at <= ?")
        values.append(date_to + "T23:59:59+99:99")
    notify = str(params.get("notify", "") or "").strip()
    if notify in ("0", "1"):
        where.append("notify_sent = ?")
        values.append(int(notify))
    return " AND ".join(where), values


def filter_options(db_path, open_id: str) -> dict[str, list[str]]:
    """返回当前用户已有记录中的筛选候选值，不跨 open_id。"""
    columns = {
        "applicant": "applicant",
        "rule": "rule_name",
        "item_type": "item_type",
        "uom": "uom",
    }
    options = {}
    with closing(_connect(db_path)) as conn, conn:
        for key, column in columns.items():
            rows = conn.execute(
                f"""SELECT DISTINCT {column} AS value
                FROM sign_actions
                WHERE open_id = ? AND trim({column}) <> ''
                ORDER BY {column} COLLATE NOCASE""",
                (open_id,),
            ).fetchall()
            options[key] = [str(row["value"]) for row in rows]
    return options


def query_actions(db_path, open_id: str, params=None, limit: int = 200, offset: int = 0):
    params = params or {}
    where, values = _filters(open_id, params)
    sql = "SELECT * FROM sign_actions WHERE %s ORDER BY created_at DESC LIMIT ? OFFSET ?" % where
    with closing(_connect(db_path)) as conn, conn:
        rows = conn.execute(sql, values + [max(1, min(limit, 100000)), max(0, offset)]).fetchall()
    return [dict(row) for row in rows]


def summarize(db_path, open_id: str, params=None):
    params = params or {}
    where, values = _filters(open_id, params)
    with closing(_connect(db_path)) as conn, conn:
        total = conn.execute("SELECT COUNT(*) FROM sign_actions WHERE %s" % where, values).fetchone()[0]
        by_action = dict(conn.execute(
            "SELECT action,COUNT(*) FROM sign_actions WHERE %s GROUP BY action" % where, values
        ).fetchall())
        by_source = dict(conn.execute(
            "SELECT source,COUNT(*) FROM sign_actions WHERE %s GROUP BY source" % where, values
        ).fetchall())
        by_day = [dict(row) for row in conn.execute(
            """SELECT substr(created_at,1,10) AS day,COUNT(*) AS count
            FROM sign_actions WHERE %s GROUP BY day ORDER BY day DESC LIMIT 30""" % where, values
        ).fetchall()]
    return {"total": total, "by_action": by_action, "by_source": by_source, "by_day": by_day}


def export_excel(db_path, open_id: str, params=None) -> bytes:
    from openpyxl import Workbook
    rows = query_actions(db_path, open_id, params or {}, limit=100000)
    wb = Workbook()
    ws = wb.active
    ws.title = "我的签核记录"
    headers = ["时间", "申请单号", "申请人", "描述", "单位", "类别", "操作", "来源", "规则", "原因", "群通知", "状态"]
    ws.append(headers)
    for row in rows:
        ws.append([
            row["created_at"], row["application_no"], row["applicant"], row["description"],
            row["uom"], row["item_type"], "签核" if row["action"] == "approve" else "拒签",
            "自动" if row["source"] == "auto" else "手动", row["rule_name"], row["reason"],
            "已发送" if row["notify_sent"] else "未发送", row["status"],
        ])
    ws.freeze_panes = "A2"
    widths = [22, 14, 14, 50, 10, 10, 10, 10, 24, 36, 10, 12]
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + index)].width = width
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()
